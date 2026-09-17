"""Report, storage and email-drafting tests. All offline, all in tmp_path."""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

import config
from src import emailer, reporter, storage
from src.extractor import extract_certificate, field_confidences, field_values
from src.matcher import match_vendor
from src.validator import overall_status, validate

TEMPLATES_DIR = Path(__file__).resolve().parents[1] / "templates"
EXPECTED_TABS = [
    "Summary",
    "All Vendors",
    "Expiring",
    "Non-Compliant",
    "Needs Review",
    "Raw Extractions",
]


@pytest.fixture(scope="module")
def audited_db(tmp_path_factory, fixture_pdfs) -> Path:
    """Run the real pipeline over the fixture PDFs into a throwaway database."""
    from src.sample_generator import sample_specs, write_vendors_csv

    workdir = tmp_path_factory.mktemp("audit")
    vendors_csv = workdir / "vendors.csv"
    write_vendors_csv(sample_specs(), vendors_csv)

    from src.matcher import load_vendors

    vendors = load_vendors(vendors_csv)
    conn = storage.connect(workdir / "coi.db")
    storage.upsert_vendors(conn, vendors)

    today = date.today()
    for pdf in sorted(fixture_pdfs.values()):
        extraction = extract_certificate(pdf)
        match = match_vendor(extraction["fields"]["insured"]["value"], vendors)
        findings = validate(
            field_values(extraction),
            field_confidences(extraction),
            config.requirements_for(match["category"]),
            today,
        )
        cert_id = storage.save_certificate(conn, extraction, match, overall_status(findings))
        storage.replace_findings(conn, cert_id, findings)
    conn.close()
    return workdir / "coi.db"


@pytest.fixture
def conn(audited_db):
    connection = storage.connect(audited_db)
    yield connection
    connection.close()


# --------------------------------------------------------------------------
# storage
# --------------------------------------------------------------------------

def test_all_eight_certificates_are_stored(conn):
    assert len(storage.all_certificates(conn)) == 8


def test_duplicate_hash_is_detected(conn):
    first = storage.all_certificates(conn)[0]
    assert storage.find_by_hash(conn, first["file_hash"]) == first["id"]
    assert storage.find_by_hash(conn, "nope") is None


def test_current_certificate_is_one_per_vendor(conn):
    current = storage.current_certificates(conn)
    names = [row["vendor_name"] for row in current]
    assert len(names) == len(set(names)) == 8


def test_update_field_bumps_confidence_and_revalidates(conn):
    row = next(r for r in storage.all_certificates(conn) if r["gl_each_occurrence"] == 500_000)
    storage.update_field(conn, row["id"], "gl_each_occurrence", 2_000_000)

    updated = storage.get_certificate(conn, row["id"])
    assert updated["gl_each_occurrence"] == 2_000_000
    assert storage.row_confidences(updated)["gl_each_occurrence"] == 1.0

    findings = validate(
        storage.row_to_fields(updated),
        storage.row_confidences(updated),
        config.requirements_for(updated["category"]),
        date.today(),
    )
    assert overall_status(findings) == "PASS"

    storage.update_field(conn, row["id"], "gl_each_occurrence", 500_000)


def test_row_to_fields_round_trips_types(conn):
    row = storage.all_certificates(conn)[0]
    fields = storage.row_to_fields(row)
    assert isinstance(fields["policy_exp"], date)
    assert isinstance(fields["gl_each_occurrence"], int)
    assert isinstance(fields["addl_insd"], bool)


# --------------------------------------------------------------------------
# report
# --------------------------------------------------------------------------

def test_workbook_is_written_with_every_tab(conn, tmp_path):
    out_path = reporter.write_report(conn, tmp_path, date.today())
    assert out_path.exists()
    assert out_path.name == f"coi_report_{date.today().isoformat()}.xlsx"

    workbook = openpyxl.load_workbook(out_path)
    assert workbook.sheetnames == EXPECTED_TABS


def test_summary_tab_reports_counts_and_disclaimer(conn, tmp_path):
    out_path = reporter.write_report(conn, tmp_path, date.today())
    sheet = openpyxl.load_workbook(out_path)["Summary"]
    values = {row[0]: row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}

    assert values["Vendors with a certificate on file"] == 8
    assert values["PASS"] == 2
    assert values["WARN"] == 2
    assert values["FAIL"] == 4
    assert values["Expired"] == 1
    assert values["Expiring within 30 days"] == 2
    assert values["Compliance %"] == "25.0%"
    assert values["Note"] == reporter.DISCLAIMER


def test_all_vendors_tab_has_a_row_per_vendor_and_is_frozen(conn, tmp_path):
    sheet = openpyxl.load_workbook(reporter.write_report(conn, tmp_path, date.today()))["All Vendors"]
    assert sheet.max_row == 9  # header + 8 vendors
    assert sheet.freeze_panes == "A2"
    headers = [cell.value for cell in sheet[1]]
    for column in ("vendor_name", "status", "earliest_expiry", "days_remaining", "findings"):
        assert column in headers


def test_status_cells_are_colour_filled(conn, tmp_path):
    sheet = openpyxl.load_workbook(reporter.write_report(conn, tmp_path, date.today()))["All Vendors"]
    headers = [cell.value for cell in sheet[1]]
    status_col = headers.index("status") + 1
    seen = set()
    for row_index in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row_index, column=status_col)
        seen.add(cell.value)
        assert cell.fill.start_color.rgb == reporter.STATUS_FILLS[cell.value].start_color.rgb
    assert seen == {"PASS", "WARN", "FAIL"}


def test_expiring_and_noncompliant_tabs_hold_the_right_rows(conn, tmp_path):
    workbook = openpyxl.load_workbook(reporter.write_report(conn, tmp_path, date.today()))
    assert workbook["Expiring"].max_row == 3  # header + 2 expiring vendors
    assert workbook["Non-Compliant"].max_row == 5  # header + 4 failing vendors


def test_clean_run_leaves_needs_review_empty(conn, tmp_path):
    sheet = openpyxl.load_workbook(reporter.write_report(conn, tmp_path, date.today()))["Needs Review"]
    assert sheet["A1"].value == "No rows for this tab."


def test_raw_extractions_carries_values_and_confidences(conn, tmp_path):
    sheet = openpyxl.load_workbook(reporter.write_report(conn, tmp_path, date.today()))["Raw Extractions"]
    headers = [cell.value for cell in sheet[1]]
    assert "gl_each_occurrence" in headers
    assert "gl_each_occurrence__conf" in headers
    assert sheet.max_row == 9


def test_report_on_an_empty_database_still_writes_every_tab(tmp_path):
    empty = storage.connect(tmp_path / "empty.db")
    out_path = reporter.write_report(empty, tmp_path, date.today())
    empty.close()
    assert openpyxl.load_workbook(out_path).sheetnames == EXPECTED_TABS


# --------------------------------------------------------------------------
# emails
# --------------------------------------------------------------------------

def test_drafts_one_email_per_actionable_vendor(conn, tmp_path):
    drafts = emailer.draft_emails(
        conn, TEMPLATES_DIR, tmp_path / "emails", "Northwind Property Management", "Risk Management"
    )
    assert len(drafts) == 6  # 2 WARN + 4 FAIL
    assert all(path.suffix == ".eml" for path in drafts)
    assert (tmp_path / "emails" / "mail_merge.csv").exists()


def test_draft_contains_headers_findings_and_no_send(conn, tmp_path):
    emailer.draft_emails(
        conn, TEMPLATES_DIR, tmp_path / "emails", "Northwind Property Management", "Risk Management"
    )
    path = next((tmp_path / "emails").glob("summit_hvac*.eml"))
    text = path.read_text(encoding="utf-8")

    assert "To: ops@summithvac.example.com" in text
    assert "Subject:" in text
    assert "draft-only; not sent" in text
    assert "Summit HVAC Solutions LLC" in text
    assert "Northwind Property Management" in text.replace("=\n", "")


def test_template_choice_follows_the_findings():
    expiring = [{"code": "EXPIRING_30", "severity": "WARN", "message": ""}]
    deficient = [{"code": "WC_MISSING", "severity": "FAIL", "message": ""}]
    expired = [{"code": "EXPIRED", "severity": "FAIL", "message": ""}]

    assert emailer.choose_template(expiring, TEMPLATES_DIR).name == "expiring.txt"
    assert emailer.choose_template(deficient, TEMPLATES_DIR).name == "deficient.txt"
    assert emailer.choose_template(expired, TEMPLATES_DIR).name == "deficient.txt"


def test_mail_merge_csv_lists_every_draft(conn, tmp_path):
    import csv

    emailer.draft_emails(
        conn, TEMPLATES_DIR, tmp_path / "emails", "Northwind Property Management", "Risk Management"
    )
    with (tmp_path / "emails" / "mail_merge.csv").open(encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 6
    assert all(row["to"] and row["subject"] and row["eml_file"] for row in rows)
    assert {row["status"] for row in rows} == {"WARN", "FAIL"}


def test_templates_render_without_leftover_placeholders(conn, tmp_path):
    emailer.draft_emails(
        conn, TEMPLATES_DIR, tmp_path / "emails", "Northwind Property Management", "Risk Management"
    )
    for path in (tmp_path / "emails").glob("*.eml"):
        body = path.read_text(encoding="utf-8").replace("=\n", "")
        assert "{" not in body and "}" not in body


def test_expiry_date_is_rendered_us_style(conn, tmp_path):
    emailer.draft_emails(
        conn, TEMPLATES_DIR, tmp_path / "emails", "Northwind Property Management", "Risk Management"
    )
    expected = (date.today() + timedelta(days=20)).strftime("%m/%d/%Y")
    text = next((tmp_path / "emails").glob("summit_hvac*.eml")).read_text(encoding="utf-8")
    assert expected in text.replace("=\n", "")
